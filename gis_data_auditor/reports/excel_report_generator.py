"""
Generador de Reportes Excel profesional utilizando openpyxl.
Crea 5 pestañas formateadas: Resumen, Datasets, Detalle_Campos, Incidencias y Estandar_Utilizado.
"""

import logging
from pathlib import Path
from typing import Union
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from gis_data_auditor.models.report import AuditReport
from gis_data_auditor.models.validation import Severity

logger = logging.getLogger(__name__)


class ExcelReportGenerator:
    """Genera un archivo Excel (.xlsx) multidescripción y profesional a partir de AuditReport."""

    def __init__(self):
        # Estilos corporativos
        self.header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")  # Azul marino
        self.header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        
        self.sub_header_fill = PatternFill(start_color="4A777A", end_color="4A777A", fill_type="solid")
        self.sub_header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

        self.title_font = Font(name="Segoe UI", size=14, bold=True, color="1B365D")
        self.normal_font = Font(name="Segoe UI", size=10)
        self.bold_font = Font(name="Segoe UI", size=10, bold=True)

        self.fill_error = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")     # Rojo claro
        self.fill_warning = PatternFill(start_color="FCF3CF", end_color="FCF3CF", fill_type="solid")   # Amarillo claro
        self.fill_success = PatternFill(start_color="D4EFDF", end_color="D4EFDF", fill_type="solid")   # Verde claro
        self.fill_info = PatternFill(start_color="EBF5FB", end_color="EBF5FB", fill_type="solid")      # Azul claro

        thin_side = Side(border_style="thin", color="D5D8DC")
        self.thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    def generate_report(self, report: AuditReport, output_path: Union[str, Path]) -> Path:
        """Construye el libro Excel y lo guarda en output_path."""
        out_file = Path(output_path)
        out_file.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        # Eliminar hoja por defecto
        wb.remove(wb.active)

        self._build_sheet_resumen(wb, report)
        self._build_sheet_datasets(wb, report)
        self._build_sheet_detalle_campos(wb, report)
        self._build_sheet_incidencias(wb, report)
        self._build_sheet_estandar(wb, report)

        wb.save(out_file)
        logger.info(f"Reporte Excel generado exitosamente en '{out_file}'")
        return out_file

    def _build_sheet_resumen(self, wb: openpyxl.Workbook, report: AuditReport) -> None:
        ws = wb.create_sheet(title="Resumen")
        ws.views.sheetView[0].showGridLines = True

        # Título
        ws.cell(row=1, column=1, value="REPORTE CONSOLIDADO DE AUDITORÍA GIS").font = self.title_font
        ws.cell(row=2, column=1, value=f"Fecha y Hora: {report.timestamp.strftime('%Y-%m-%d %H:%M:%S')}").font = self.normal_font
        ws.cell(row=3, column=1, value=f"Carpeta Raíz: {report.root_folder}").font = self.normal_font
        ws.cell(row=4, column=1, value=f"Archivo Estándar: {report.standard_file}").font = self.normal_font

        # Tabla 1: Métricas de Archivos
        ws.cell(row=6, column=1, value="Métrica de Archivos").font = self.sub_header_font
        ws.cell(row=6, column=1).fill = self.sub_header_fill
        ws.cell(row=6, column=2, value="Cantidad").font = self.sub_header_font
        ws.cell(row=6, column=2).fill = self.sub_header_fill

        metrics = [
            ("Total Shapefiles Encontrados", report.total_shps_found),
            ("Total Shapefiles Leídos Exitosamente", report.total_shps_read),
            ("Total Shapefiles Auditados con Esquema", report.total_shps_audited),
            ("Sin Esquema Coincidente en Excel", report.total_no_match),
            ("Con Coincidencia Ambigua en Excel", report.total_ambiguous_match),
            ("Archivos Conformes (Sin Errores ni Advertencias)", report.total_clean_files),
            ("Archivos con Errores", report.total_files_with_errors),
            ("Archivos solo con Advertencias", report.total_files_with_warnings),
        ]

        for i, (label, val) in enumerate(metrics, start=7):
            c1 = ws.cell(row=i, column=1, value=label)
            c2 = ws.cell(row=i, column=2, value=val)
            c1.font = self.normal_font
            c2.font = self.bold_font
            c1.border = self.thin_border
            c2.border = self.thin_border
            
            if "Errores" in label and val > 0:
                c2.fill = self.fill_error
            elif "Advertencias" in label and val > 0:
                c2.fill = self.fill_warning
            elif "Conformes" in label and val > 0:
                c2.fill = self.fill_success

        # Tabla 2: Incidencias por Severidad
        ws.cell(row=17, column=1, value="Severidad de Incidencias").font = self.sub_header_font
        ws.cell(row=17, column=1).fill = self.sub_header_fill
        ws.cell(row=17, column=2, value="Total Incidencias").font = self.sub_header_font
        ws.cell(row=17, column=2).fill = self.sub_header_fill

        err_cnt = sum(1 for iss in report.issues if iss.severity == Severity.ERROR)
        warn_cnt = sum(1 for iss in report.issues if iss.severity == Severity.WARNING)
        info_cnt = sum(1 for iss in report.issues if iss.severity == Severity.INFO)

        sev_rows = [
            ("ERROR (Crítico)", err_cnt, self.fill_error),
            ("WARNING (Advertencia)", warn_cnt, self.fill_warning),
            ("INFO (Informativo)", info_cnt, self.fill_info)
        ]

        for idx, (sev_name, count, fill) in enumerate(sev_rows, start=18):
            c1 = ws.cell(row=idx, column=1, value=sev_name)
            c2 = ws.cell(row=idx, column=2, value=count)
            c1.font = self.normal_font
            c2.font = self.bold_font
            c1.border = self.thin_border
            c2.border = self.thin_border
            if count > 0:
                c2.fill = fill

        self._auto_fit_columns(ws)

    def _build_sheet_datasets(self, wb: openpyxl.Workbook, report: AuditReport) -> None:
        ws = wb.create_sheet(title="Datasets")
        headers = [
            "Nombre", "Ruta Completa", "Esquema Esperado", "Estado Asociación",
            "Estado Lectura", "Geometría", "CRS", "Registros", "Total Campos",
            "Errores", "Advertencias", "Estado Final"
        ]
        ws.append(headers)
        self._format_header_row(ws)

        for ds in report.datasets:
            ds_issues = [i for i in report.issues if i.dataset_name == ds.name]
            err_count = sum(1 for i in ds_issues if i.severity == Severity.ERROR)
            warn_count = sum(1 for i in ds_issues if i.severity == Severity.WARNING)

            if err_count > 0 or ds.read_status == "READ_ERROR" or ds.match_status == "AMBIGUOUS_MATCH":
                final_status = "ERROR"
            elif warn_count > 0 or ds.match_status == "NO_MATCH":
                final_status = "ADVERTENCIA"
            else:
                final_status = "CONFORME"

            row = [
                ds.name,
                str(ds.filepath),
                ds.associated_schema or "N/A",
                ds.match_status,
                ds.read_status,
                ds.geometry_type,
                ds.crs_name,
                ds.feature_count,
                len(ds.fields),
                err_count,
                warn_count,
                final_status
            ]
            ws.append(row)
            
            # Formato de la última fila agregada
            r_idx = ws.max_row
            status_cell = ws.cell(row=r_idx, column=12)
            if final_status == "ERROR":
                status_cell.fill = self.fill_error
            elif final_status == "ADVERTENCIA":
                status_cell.fill = self.fill_warning
            elif final_status == "CONFORME":
                status_cell.fill = self.fill_success

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        self._auto_fit_columns(ws)

    def _build_sheet_detalle_campos(self, wb: openpyxl.Workbook, report: AuditReport) -> None:
        ws = wb.create_sheet(title="Detalle_Campos")
        headers = [
            "Dataset", "Esquema", "Campo", "Estado Campo", "Tipo Esperado", "Tipo Real",
            "Longitud Esperada", "Longitud Real", "Decimales Esperados", "Escala Real", "Resultado Campo"
        ]
        ws.append(headers)
        self._format_header_row(ws)

        for ds in report.datasets:
            if ds.match_status != "MATCH" or ds.read_status != "SUCCESS":
                continue

            std_def = report.standards_used.get(ds.associated_schema.upper()) if ds.associated_schema else None
            if not std_def:
                continue

            all_field_names = set(std_def.fields.keys()) | set(ds.fields.keys())

            for f_upper in sorted(all_field_names):
                f_std = std_def.get_field(f_upper)
                f_real = ds.get_field(f_upper)

                if f_std and f_real:
                    status_campo = "EXISTENTE"
                elif f_std and not f_real:
                    status_campo = "FALTANTE"
                else:
                    status_campo = "ADICIONAL"

                # Obtener problemas específicos de este campo
                field_issues = [
                    i for i in report.issues
                    if i.dataset_name == ds.name and i.field_name.upper() == f_upper
                ]

                has_err = any(i.severity == Severity.ERROR for i in field_issues)
                has_warn = any(i.severity == Severity.WARNING for i in field_issues)

                if has_err or status_campo == "FALTANTE":
                    res_campo = "ERROR"
                elif has_warn or status_campo == "ADICIONAL":
                    res_campo = "ADVERTENCIA"
                else:
                    res_campo = "CONFORME"

                exp_len_str = str(f_std.expected_length) if (f_std and f_std.has_length_spec) else "NO_ESPECIFICADO"
                exp_dec_str = str(f_std.expected_decimals) if (f_std and f_std.has_decimals_spec) else "NO_ESPECIFICADO"

                row = [
                    ds.name,
                    std_def.schema_id,
                    f_std.name if f_std else f_real.name,
                    status_campo,
                    f_std.expected_type if f_std else "N/A",
                    f_real.type_name if f_real else "N/A",
                    exp_len_str,
                    f_real.length if (f_real and f_real.length is not None) else "N/A",
                    exp_dec_str,
                    f_real.scale if (f_real and f_real.scale is not None) else "N/A",
                    res_campo
                ]
                ws.append(row)
                
                r_idx = ws.max_row
                res_cell = ws.cell(row=r_idx, column=11)
                if res_campo == "ERROR":
                    res_cell.fill = self.fill_error
                elif res_campo == "ADVERTENCIA":
                    res_cell.fill = self.fill_warning
                elif res_campo == "CONFORME":
                    res_cell.fill = self.fill_success

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        self._auto_fit_columns(ws)

    def _build_sheet_incidencias(self, wb: openpyxl.Workbook, report: AuditReport) -> None:
        ws = wb.create_sheet(title="Incidencias")
        headers = [
            "ID Incidencia", "Archivo SHP", "Ruta Completa", "Esquema", "Campo",
            "Regla", "Categoría", "Severidad", "Valor Esperado", "Valor Encontrado",
            "Descripción de la Incidencia", "Recomendación"
        ]
        ws.append(headers)
        self._format_header_row(ws)

        for iss in report.issues:
            row = [
                iss.issue_id,
                iss.dataset_name,
                iss.filepath,
                iss.schema_id,
                iss.field_name,
                iss.rule_name,
                iss.category.value,
                iss.severity.value,
                str(iss.expected_value),
                str(iss.found_value),
                iss.message,
                iss.recommendation
            ]
            ws.append(row)

            r_idx = ws.max_row
            sev_cell = ws.cell(row=r_idx, column=8)
            if iss.severity == Severity.ERROR:
                sev_cell.fill = self.fill_error
            elif iss.severity == Severity.WARNING:
                sev_cell.fill = self.fill_warning
            elif iss.severity == Severity.INFO:
                sev_cell.fill = self.fill_info

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        self._auto_fit_columns(ws)

    def _build_sheet_estandar(self, wb: openpyxl.Workbook, report: AuditReport) -> None:
        ws = wb.create_sheet(title="Estandar_Utilizado")
        headers = ["Esquema (shp)", "Nombre Campo", "Tipo Esperado", "Longitud Esperada", "Decimales Esperados", "Fila Origen Excel"]
        ws.append(headers)
        self._format_header_row(ws)

        for schema_id, std_def in report.standards_used.items():
            for f_def in std_def.fields.values():
                len_str = str(f_def.expected_length) if f_def.has_length_spec else "NA"
                dec_str = str(f_def.expected_decimals) if f_def.has_decimals_spec else "NA"
                row = [
                    std_def.schema_id,
                    f_def.name,
                    f_def.expected_type,
                    len_str,
                    dec_str,
                    f_def.row_number
                ]
                ws.append(row)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        self._auto_fit_columns(ws)

    def _format_header_row(self, ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
        for cell in ws[1]:
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _auto_fit_columns(self, ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)
